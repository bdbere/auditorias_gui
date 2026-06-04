import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import matplotlib.pyplot as plt
import textwrap 

# --- Configuración de Matplotlib ---
plt.style.use('ggplot')

# --- Estructuras de Datos Principales ---

escala = {
    "Excelente": 5, "Bueno": 4, "Regular": 3, "Malo": 2, "Muy Malo": 1,
    "excelente": 5, "bueno": 4, "regular": 3, "malo": 2, "muy malo": 1
}

# --- ACTUALIZACIÓN: Nuevos rangos de columnas solicitados ---
servicios_columnas = {
    "Centro de información": ("H", "O"),
    "Coordinación": ("R", "Y"),
    "Titulación": ("AB", "AJ"),
    "Recursos Financieros": ("AL", "AQ"),
    "Residencias": ("AS", "AZ"),
    "Laboratorio de cómputo": ("BA", "BE"),
    "Servicio Social": ("BG", "BM"),
    "Servicios Escolares": ("BO", "BW"),
    "Cafetería": ("BY", "CE"),
    "Centro de Copiado": ("CG", "CM"),
    "Seguridad": ("CO", "CS"),
    "Limpieza": ("CU", "CZ")
}

# Variables globales
df_datos = None
nombre_archivo = ""

# --- Funciones Auxiliares ---

def col_letra_a_num(letra):
    """Convierte letras de Excel (ej. 'AB') a índice numérico base 0."""
    num = 0
    for c in letra:
        num = num * 26 + (ord(c.upper()) - ord('A') + 1)
    return num - 1

def aplicar_estilos_celda(ws, fila, col, valor, negrita=False, alineacion="center", color_fondo=None):
    """Ayuda para formato de Excel."""
    celda = ws.cell(row=fila, column=col, value=valor)
    if negrita: celda.font = Font(bold=True)
    celda.alignment = Alignment(horizontal=alineacion, vertical="center", wrap_text=True)
    celda.border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    if color_fondo: celda.fill = PatternFill(start_color=color_fondo, end_color=color_fondo, fill_type="solid")
    return celda

# --- Lógica de Gráficos ---

def calcular_promedios_servicio(df_carrera, col_inicio, col_fin):
    """Calcula los promedios de un rango de columnas para graficar."""
    idx_inicio = col_letra_a_num(col_inicio)
    idx_fin = col_letra_a_num(col_fin)
    
    # Validación de índices para evitar errores si el archivo es más corto de lo esperado
    if idx_fin >= len(df_carrera.columns):
        return [], []

    cols_servicio = df_carrera.columns[idx_inicio : idx_fin + 1]
    
    preguntas = []
    promedios = []
    
    for pregunta in cols_servicio:
        datos = df_carrera[pregunta].dropna().astype(str).str.strip().str.title()
        valores_numericos = datos.map(escala)
        
        if not valores_numericos.isna().all():
            prom = valores_numericos.mean()
            preguntas.append(pregunta)
            promedios.append(prom)
    
    return preguntas, promedios

def mostrar_ventana_graficos():
    global df_datos
    if df_datos is None: return

    ventana_graf = tk.Toplevel()
    ventana_graf.title("Generador de Gráficos")
    ventana_graf.geometry("450x350")
    
    ttk.Label(ventana_graf, text="Selecciona Carrera y Servicio:", font=("Arial", 12, "bold")).pack(pady=10)
    
    carreras = sorted(df_datos['Carrera'].dropna().unique().astype(str))
    ttk.Label(ventana_graf, text="Carrera:").pack(pady=2)
    combo_carrera = ttk.Combobox(ventana_graf, values=carreras, width=50, state="readonly")
    combo_carrera.pack(pady=5)
    if carreras: combo_carrera.current(0)
    
    servicios = list(servicios_columnas.keys())
    ttk.Label(ventana_graf, text="Servicio:").pack(pady=2)
    combo_servicio = ttk.Combobox(ventana_graf, values=servicios, width=50, state="readonly")
    combo_servicio.pack(pady=5)
    combo_servicio.current(0)
    
    def generar_plot():
        carrera_sel = combo_carrera.get()
        servicio_sel = combo_servicio.get()
        
        if not carrera_sel or not servicio_sel: return
        
        df_carrera = df_datos[df_datos['Carrera'] == carrera_sel]
        rango = servicios_columnas[servicio_sel]
        
        preguntas, promedios = calcular_promedios_servicio(df_carrera, rango[0], rango[1])
        
        if not preguntas:
            messagebox.showinfo("Sin datos", f"No hay respuestas válidas o columnas para {servicio_sel} en {carrera_sel}.")
            return
            
        plt.figure(figsize=(12, 7)) # Un poco más grande para ver mejor
        etiquetas_x = [textwrap.fill(str(p), width=25) for p in preguntas]
        
        barras = plt.bar(etiquetas_x, promedios, color='#4F81BD')
        
        plt.title(f"{servicio_sel}\n({carrera_sel})", fontsize=14)
        plt.ylabel("Promedio (1-5)", fontsize=12)
        plt.ylim(0, 5.5)
        
        for bar in barras:
            height = bar.get_height()
            plt.text(bar.get_x() + bar.get_width()/2., height,
                     f'{height:.2f}',
                     ha='center', va='bottom', fontsize=10, fontweight='bold')
        
        plt.xticks(rotation=45, ha='right', fontsize=9)
        plt.tight_layout()
        plt.show()

    ttk.Button(ventana_graf, text="Ver Gráfico", command=generar_plot).pack(pady=20, ipadx=10, ipady=5)


# --- Lógica de Excel ---

def procesar_y_generar_excel():
    global df_datos
    if df_datos is None:
        messagebox.showwarning("Atención", "Primero debes cargar un archivo.")
        return

    ruta_guardar = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Archivos de Excel", "*.xlsx")],
        title="Guardar Reporte de Resultados"
    )
    if not ruta_guardar: return

    try:
        wb = Workbook()
        if "Sheet" in wb.sheetnames: del wb["Sheet"]
        
        carreras_unicas = df_datos['Carrera'].dropna().unique()
        
        for carrera in carreras_unicas:
            # Limpieza nombre hoja (max 31 chars)
            nombre_hoja = str(carrera)[:30].replace("/", "-").replace(":", "")
            ws = wb.create_sheet(title=nombre_hoja)
            df_carrera = df_datos[df_datos['Carrera'] == carrera]
            fila_actual = 2
            
            ws.merge_cells(start_row=fila_actual, start_column=2, end_row=fila_actual, end_column=10)
            aplicar_estilos_celda(ws, fila_actual, 2, f"Resultados para: {carrera}", negrita=True, color_fondo="D9EAD3").font = Font(size=14, bold=True)
            fila_actual += 2
            
            for nombre_servicio, (col_inicio, col_fin) in servicios_columnas.items():
                encabezados = ["Excelente (D)", "Bueno (E)", "Regular (F)", "Malo (G)", "Muy Malo (H)", "Promedio (I)", "Total (J)"]
                
                aplicar_estilos_celda(ws, fila_actual, 2, nombre_servicio, negrita=True, alineacion="left", color_fondo="EFEFEF")
                for i, texto in zip(range(4, 11), encabezados):
                    aplicar_estilos_celda(ws, fila_actual, i, texto, negrita=True, color_fondo="CCE5FF")
                
                fila_actual += 1
                idx_inicio = col_letra_a_num(col_inicio)
                idx_fin = col_letra_a_num(col_fin)
                
                # Verificación de rango por si el excel es más corto
                if idx_fin >= len(df_datos.columns):
                    continue

                cols_servicio = df_datos.columns[idx_inicio : idx_fin + 1]
                
                sumas_promedios = 0
                count_preguntas_validas = 0
                
                for pregunta in cols_servicio:
                    datos_pregunta = df_carrera[pregunta]
                    respuestas = datos_pregunta.dropna().astype(str).str.strip().str.title()
                    conteo = respuestas.value_counts()
                    
                    n_exc = conteo.get("Excelente", 0)
                    n_bue = conteo.get("Bueno", 0)
                    n_reg = conteo.get("Regular", 0)
                    n_mal = conteo.get("Malo", 0)
                    n_muy_mal = conteo.get("Muy Malo", 0) + conteo.get("Muy malo", 0)
                    
                    total = n_exc + n_bue + n_reg + n_mal + n_muy_mal
                    
                    if total > 0:
                        promedio = ((n_exc*5)+(n_bue*4)+(n_reg*3)+(n_mal*2)+(n_muy_mal*1))/total
                        sumas_promedios += promedio
                        count_preguntas_validas += 1
                    else:
                        promedio = 0
                    
                    ws.column_dimensions['C'].width = 50
                    aplicar_estilos_celda(ws, fila_actual, 3, pregunta, alineacion="left")
                    for idx, val in enumerate([n_exc, n_bue, n_reg, n_mal, n_muy_mal]):
                        aplicar_estilos_celda(ws, fila_actual, 4 + idx, val)
                    
                    celda_prom = aplicar_estilos_celda(ws, fila_actual, 9, f"{promedio:.2f}", negrita=True)
                    if promedio < 3.0 and total > 0: celda_prom.font = Font(color="FF0000", bold=True)
                    aplicar_estilos_celda(ws, fila_actual, 10, total)
                    fila_actual += 1
                
                promedio_bloque = sumas_promedios / count_preguntas_validas if count_preguntas_validas > 0 else 0
                ws.merge_cells(start_row=fila_actual, start_column=3, end_row=fila_actual, end_column=8)
                aplicar_estilos_celda(ws, fila_actual, 3, "Promedio General del Servicio:", negrita=True, alineacion="right", color_fondo="FFCC00")
                aplicar_estilos_celda(ws, fila_actual, 9, f"{promedio_bloque:.2f}", negrita=True, color_fondo="FFCC00")
                fila_actual += 3

        wb.save(ruta_guardar)
        messagebox.showinfo("Éxito", f"Excel generado en:\n{ruta_guardar}")
        boton_graficos.config(state=tk.NORMAL)

    except Exception as e:
        messagebox.showerror("Error", f"Error al generar reporte:\n{e}")

# --- Funciones GUI Principal ---

def cargar_archivo():
    global df_datos, nombre_archivo
    archivo_path = filedialog.askopenfilename(
        title="Seleccionar archivo", filetypes=[("Excel/CSV", "*.xlsx *.csv")]
    )
    if not archivo_path: return

    try:
        if archivo_path.endswith('.csv'): df_datos = pd.read_csv(archivo_path)
        else: df_datos = pd.read_excel(archivo_path)

        # --- ACTUALIZACIÓN: Buscar Carrera en la columna G (índice 6) ---
        if "Carrera" not in df_datos.columns:
            # Checar si existe índice 6 (Columna G)
            if len(df_datos.columns) > 6: 
                col_g = df_datos.columns[6]
                df_datos.rename(columns={col_g: 'Carrera'}, inplace=True)
            else: 
                messagebox.showerror("Error", "No se encuentra columna G (Carrera) en el archivo.")
                return
        
        nombre_archivo = archivo_path.split('/')[-1]
        label_archivo.config(text=f"Archivo: {nombre_archivo}", foreground="green")
        boton_procesar.config(state=tk.NORMAL)
        boton_graficos.config(state=tk.NORMAL) 
        messagebox.showinfo("Carga Exitosa", "Archivo cargado. Columna G asignada como Carrera.")

    except Exception as e:
        messagebox.showerror("Error", f"No se pudo leer el archivo:\n{e}")

def main():
    global label_archivo, boton_procesar, boton_graficos
    root = tk.Tk()
    root.title("Analizador de Auditorías v2.0")
    root.geometry("600x450")

    main_frame = ttk.Frame(root, padding="20")
    main_frame.pack(expand=True, fill=tk.BOTH)

    ttk.Label(main_frame, text="Sistema de Análisis de Auditorías", font=("Helvetica", 16, "bold")).pack(pady=10)

    ttk.Button(main_frame, text="1. Cargar Archivo", command=cargar_archivo).pack(pady=10, fill=tk.X)
    label_archivo = ttk.Label(main_frame, text="Sin archivo", foreground="gray")
    label_archivo.pack(pady=5)

    boton_procesar = ttk.Button(main_frame, text="2. Generar Reporte Excel", command=procesar_y_generar_excel, state=tk.DISABLED)
    boton_procesar.pack(pady=10, fill=tk.X)

    boton_graficos = ttk.Button(main_frame, text="3. Generar Gráficos", command=mostrar_ventana_graficos, state=tk.DISABLED)
    boton_graficos.pack(pady=10, fill=tk.X)

    root.mainloop()

if __name__ == "__main__":
    main()